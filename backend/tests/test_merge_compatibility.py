"""Regression checks for merging the remote pipeline with account workspaces."""
import asyncio
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook
from fastapi.testclient import TestClient

from app.application import create_app
from app.config import settings
from app.core.auth_service import COOKIE_NAME
from app.core.docx_exporter import DocxExporter
from app.core.excel_exporter import ExcelExporter
from app.core.telemetry_service import TelemetryService


class MergeCompatibilityTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.app = create_app(self.directory.name, ["http://testserver"])
        self.client = TestClient(self.app)
        self.client.__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.directory.cleanup()

    def account(self, name):
        self.client.cookies.clear()
        response = self.client.post("/api/auth/register", json={
            "username": name, "password": "Test-password-12345",
        })
        self.assertEqual(response.status_code, 201, response.text)
        cookie = {"Cookie": f"{COOKIE_NAME}={response.cookies[COOKIE_NAME]}"}
        workspace = self.app.state.workspaces.get(response.json()["user"]["id"])
        self.client.cookies.clear()
        return workspace, cookie

    @staticmethod
    def add_cell(workspace, cell_id="cell_scope"):
        workspace.cell_lake.insert_cells([{
            "cell_id": cell_id, "file_name": "school.xlsx", "sheet_name": "Summary",
            "row_idx": 2, "col_idx": 1, "cell_ref": "A2",
            "metric_path": "教师人数", "raw_value": "42", "numeric_value": 42,
        }])

    def test_container_health_check_does_not_require_account(self):
        response = self.client.get("/healthz")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"status": "healthy"})
        self.assertEqual(self.client.get("/api/health").status_code, 401)

    def test_docker_frontend_origin_can_register(self):
        with tempfile.TemporaryDirectory() as root:
            with patch.dict("os.environ", {"APP_ALLOWED_ORIGINS": "http://localhost:8080"}):
                app = create_app(root)
            with TestClient(app) as client:
                response = client.post("/api/auth/register", headers={"Origin": "http://localhost:8080"}, json={
                    "username": "docker_teacher", "password": "Test-password-12345",
                })
                self.assertEqual(response.status_code, 201, response.text)

    def test_telemetry_and_balance_are_scoped_to_account(self):
        first, first_cookie = self.account("teacher_a")
        _, second_cookie = self.account("teacher_b")
        TelemetryService.record_section_trace(
            "sec_a", "Private chapter", "OverviewSpecialist", 100, 10, 20, 0.95, 1, True,
            trace_file=first.pipeline.telemetry_file,
        )
        own = self.client.get("/api/pipeline/telemetry", headers=first_cookie)
        other = self.client.get("/api/pipeline/telemetry", headers=second_cookie)
        self.assertEqual(own.status_code, 200, own.text)
        self.assertEqual(other.status_code, 200, other.text)
        self.assertEqual(own.json()["total_sections_traced"], 1)
        self.assertEqual(other.json()["total_sections_traced"], 0)
        with patch.object(settings, "DEEPSEEK_API_KEY", "global-key-must-not-be-used"):
            balance = self.client.get("/api/pipeline/balance", headers=second_cookie)
        self.assertEqual(balance.status_code, 200, balance.text)
        self.assertFalse(balance.json()["is_configured"])
        self.assertEqual(balance.json()["token_burn_summary"]["total_sections_traced"], 0)

    def test_sse_keeps_resume_and_account_operation_lock(self):
        workspace, cookie = self.account("teacher_sse")
        self.add_cell(workspace)
        calls = []

        async def execute(**kwargs):
            calls.append(kwargs)
            yield {"type": "stage_update", "stage_id": 1, "status": "running"}

        with patch.object(workspace.pipeline, "execute_pipeline", execute):
            response = self.client.get("/api/pipeline/stream?resume=true&school_name=Test", headers=cookie)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["x-accel-buffering"], "no")
        self.assertEqual(calls, [{"school_name": "Test", "resume": True, "use_existing_outline": False}])
        self.assertFalse(workspace.operation_lock.locked())

    def test_resume_uses_account_checkpoint_and_rejects_changed_data(self):
        first, _ = self.account("teacher_resume")
        second, _ = self.account("teacher_empty")
        self.add_cell(first)
        pipe = first.pipeline
        content = "学校现有专任教师[42人][^cell_scope]。" * 12
        writes = []
        outline = [{
            "id": "sec_scope", "chapter_title": "第一章 学校概况",
            "section_title": "1.1 师资规模", "objective": "分析教师规模",
            "file_name": "school.xlsx", "sheet_name": "Summary", "table_name": "tbl_school",
            "bound_files": ["school.xlsx"],
        }]

        async def plan(**kwargs):
            pipe.sections_plan = outline
            return outline

        async def write(*args, **kwargs):
            writes.append(1)
            yield {"type": "chunk", "text": content}
            yield {"type": "done", "full_content": content}

        async def run(resume=False):
            return [event async for event in pipe.execute_pipeline(school_name="Test University", resume=resume)]

        with patch.object(pipe, "async_plan_outline", plan), patch.object(pipe.agent_runner, "stream_write_section", write):
            asyncio.run(run())
            self.assertEqual(len(writes), 1)
            checkpoint = Path(first.checkpoint_dir) / "pipeline_checkpoint.json"
            saved = json.loads(checkpoint.read_text(encoding="utf-8"))
            self.assertEqual(saved["generated_sections"][0]["content"], content)
            self.assertEqual(saved["sections"][0]["content"], content)
            resumed = asyncio.run(run(resume=True))
            self.assertEqual(len(writes), 1)
            self.assertTrue(any("断点续生" in event.get("text", "") for event in resumed))
            # 页面已经生成过的异常小节也属于“已有正文”，只允许用户显式单节重试。
            saved["generated_sections"][0]["passed"] = False
            saved["sections"][0]["passed"] = False
            checkpoint.write_text(json.dumps(saved, ensure_ascii=False), encoding="utf-8")
            resumed_failed = asyncio.run(run(resume=True))
            self.assertEqual(len(writes), 1)
            self.assertTrue(any("断点续生" in event.get("text", "") for event in resumed_failed))
            self.assertFalse((second.checkpoint_dir / "pipeline_checkpoint.json").exists())
            self.assertFalse(Path(second.pipeline.telemetry_file).exists())
            first.mark_data_changed()
            asyncio.run(run(resume=True))
            self.assertEqual(len(writes), 2)

    def test_demo_exports_are_clearly_labeled(self):
        with tempfile.TemporaryDirectory() as output:
            docx_path = DocxExporter(output).export_report(
                report_title="示例报告：海南师范大学本科教育教学质量发展检验报告",
                school_name="海南师范大学",
                sections=[{"title": "1.1 示例小节", "level": 2, "content": "示例正文", "audit": {"citations": []}}],
                charts_dir=output,
                quality_summary={"sections_total": 1, "sections_passed": 1, "sections_failed": 0, "can_deliver": True},
                is_demo=True,
            )
            text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
            self.assertIn("示例报告", text)
            self.assertIn("不得作为正式评估材料", text)

            xlsx_path = ExcelExporter(output).export_audit_workbook(
                school_name="海南师范大学",
                coverage_data=[],
                reconciliation_data=[],
                quality_summary={"sections_total": 1, "sections_passed": 1, "sections_failed": 0, "can_deliver": True},
                is_demo=True,
            )
            self.assertTrue(Path(xlsx_path).name.startswith("示例报告_"))
            workbook = load_workbook(xlsx_path, data_only=True)
            quality_rows = [
                tuple(cell.value for cell in row)
                for row in workbook["质量说明"].iter_rows(min_row=1, max_row=12, max_col=4)
            ]
            self.assertTrue(any("示例报告" in str(value) for row in quality_rows for value in row))

    def test_docx_cover_and_markdown_table_are_delivery_ready(self):
        workspace, _ = self.account("teacher_docx")
        destination = workspace.pipeline.docx_exporter.export_report(
            "Report",
            "海南师范大学",
            [{
                "title": "2.1 测试章节",
                "level": 2,
                "content": "正文说明。\n\n| 单元格 | 单位号 |\n|---|---|\n| B2 | 3100 |\n\n结尾说明。",
                "audit": {"citations": []},
            }],
            str(workspace.chart_dir),
            quality_summary={
                "needs_attention": True,
                "can_deliver": False,
                "outline_approved": False,
                "outline_score": 0.82,
                "sections_failed": 1,
                "sections_total": 1,
            },
            organization="海南师范大学",
        )
        doc = Document(destination)
        paragraphs = [paragraph.text for paragraph in doc.paragraphs]
        text = "\n".join(paragraphs + [
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        ])
        today = datetime.now()
        self.assertIn(f"{today.year} 年 {today.month} 月 {today.day} 日", text)
        self.assertIn("暂不建议正式交付", text)
        self.assertNotIn("Jev", text)
        self.assertNotIn("Cell Lake", text)
        self.assertFalse(any(line.strip().startswith("|") for line in paragraphs))
        self.assertTrue(any(len(table.columns) == 2 for table in doc.tables))

    def test_audit_workbook_has_quality_sheet_and_numeric_scope_note(self):
        workspace, _ = self.account("teacher_excel")
        path = workspace.pipeline.excel_exporter.export_audit_workbook(
            "海南师范大学",
            coverage_data=[],
            reconciliation_data=[{
                "section_title": "2.1 测试章节",
                "cited_text": "3100",
                "metric_path": "单位号",
                "file_name": "test.xls",
                "sheet_name": "Sheet1",
                "cell_ref": "B2",
                "raw_value": "3100",
                "status": "100% 吻合",
            }],
            quality_summary={
                "needs_attention": True,
                "can_deliver": False,
                "outline_approved": False,
                "outline_score": 0.82,
                "sections_total": 1,
                "sections_passed": 1,
                "sections_failed": 0,
                "failed_sections": [],
            },
        )
        wb = load_workbook(path, data_only=True)
        self.assertEqual(wb.sheetnames[0], "质量说明")
        quality = wb["质量说明"]
        self.assertEqual(quality["B8"].value, "暂不建议正式交付")
        self.assertIn("不代表语义内容全部通过", quality["A10"].value)
        self.assertIn("仅核对报告数值", wb["全文数字对账总表"]["A1"].value)

    def test_document_retains_word_toc_and_readable_citation_sources(self):
        workspace, _ = self.account("teacher_word")
        destination = workspace.pipeline.docx_exporter.export_report("Report", "Test", [{
            "title": "Chapter", "level": 2,
            "content": "教师[42人][^cell_scope]。",
            "audit": {"citations": [{"cell_id": "cell_scope", "sheet_name": "Summary", "cell_ref": "A2"}]},
        }], str(workspace.chart_dir))
        doc = Document(destination)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs] + [
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        ])
        self.assertIn("42人", text)
        self.assertIn("Summary!A2", text)
        self.assertNotIn("cell_scope", text)
        self.assertTrue(any("TOC" in node.text for node in doc.element.xpath("//w:instrText")))
        self.assertTrue(any(paragraph.style.name == "Heading 2" for paragraph in doc.paragraphs))


if __name__ == "__main__":
    unittest.main()
