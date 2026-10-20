import os
import re
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

class ExcelExporter:
    """
    《数据穿透审计与指标覆盖清单》 Excel 导出器 (.xlsx)
    生成符合高校教务处、评估专家组“查账式”免责与穿透审计要求的专业双表
    """

    def __init__(self, output_dir: str):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def export_audit_workbook(
        self,
        school_name: str,
        coverage_data: List[Dict[str, Any]],
        reconciliation_data: List[Dict[str, Any]],
        quality_summary: Optional[Dict[str, Any]] = None,
        is_demo: bool = False,
    ) -> str:
        quality_summary = quality_summary or {}
        wb = openpyxl.Workbook()

        header_fill = PatternFill(start_color="1E3D59", end_color="1E3D59", fill_type="solid")
        header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Microsoft YaHei", size=10, color="1E293B")
        status_pass_font = Font(name="Microsoft YaHei", size=10, bold=True, color="15803D")
        warning_font = Font(name="Microsoft YaHei", size=10, bold=True, color="B45309")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        warning_fill = PatternFill(start_color="FFF7E6", end_color="FFF7E6", fill_type="solid")

        quality = wb.active
        quality.title = "质量说明"
        quality.views.sheetView[0].showGridLines = True
        quality.merge_cells("A1:D1")
        title_cell = quality["A1"]
        title_cell.value = "报告质量说明"
        title_cell.fill = header_fill
        title_cell.font = Font(name="Microsoft YaHei", size=14, bold=True, color="FFFFFF")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        quality.row_dimensions[1].height = 34

        total = quality_summary.get("sections_total")
        if total is None:
            total = len({item.get("section_title") for item in reconciliation_data if item.get("section_title")})
        passed = quality_summary.get("sections_passed")
        if passed is None:
            passed = total
        failed_count = quality_summary.get("sections_failed")
        if failed_count is None:
            failed_count = 0
        outline_approved = quality_summary.get("outline_approved")
        outline_score = quality_summary.get("outline_score")
        if outline_approved is False:
            outline_text = "需复核"
            if outline_score is not None:
                outline_text += f"（{round(float(outline_score) * 100)} 分）"
        elif outline_approved is True:
            outline_text = "通过"
        else:
            outline_text = "未评估"
        if quality_summary.get("needs_attention"):
            overall = "暂不建议正式交付"
        elif quality_summary.get("can_deliver"):
            overall = "可交付"
        else:
            overall = "待核验"

        quality_rows = []
        if is_demo:
            quality_rows.append(("报告类型", "示例报告（基于示例数据，仅供功能体验）"))
        quality_rows.extend([
            ("评估学校", school_name),
            ("章节总数", total),
            ("核验通过小节", passed),
            ("核验异常小节", failed_count),
            ("大纲检查", outline_text),
            ("总体结论", overall),
        ])
        for row_idx, (label, value) in enumerate(quality_rows, 3):
            label_cell = quality.cell(row=row_idx, column=1, value=label)
            label_cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="334155")
            label_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            label_cell.border = thin_border
            label_cell.alignment = Alignment(horizontal="left", vertical="center")
            value_cell = quality.cell(row=row_idx, column=2, value=value)
            value_cell.font = warning_font if row_idx in (7, 8) and (failed_count or quality_summary.get("needs_attention")) else regular_font
            value_cell.border = thin_border
            value_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            quality.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=4)
            for col in range(3, 5):
                quality.cell(row=row_idx, column=col).border = thin_border

        quality.merge_cells("A10:D10")
        note = quality["A10"]
        note.value = "说明：本工作簿用于核对报告数值与原始单元格是否一致，不代表语义内容全部通过；语义核验结论以本页和报告封面为准。"
        note.fill = warning_fill
        note.font = Font(name="Microsoft YaHei", size=10, color="92400E")
        note.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        note.border = thin_border
        quality.row_dimensions[10].height = 38

        failed_sections = quality_summary.get("failed_sections") or []
        if failed_sections:
            failed_header_row = 12
            for col_idx, header in enumerate(["序号", "报告小节", "问题说明"], 1):
                cell = quality.cell(row=failed_header_row, column=col_idx, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
            quality.merge_cells(start_row=failed_header_row, start_column=3, end_row=failed_header_row, end_column=4)
            for idx, item in enumerate(failed_sections, 1):
                row_idx = failed_header_row + idx
                quality.cell(row=row_idx, column=1, value=idx).font = regular_font
                quality.cell(row=row_idx, column=2, value=item.get("title", "")).font = regular_font
                quality.cell(row=row_idx, column=3, value=item.get("reason", "")).font = warning_font
                quality.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=4)
                for col_idx in range(1, 5):
                    cell = quality.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    if idx % 2 == 0:
                        cell.fill = zebra_fill

        quality.column_dimensions["A"].width = 18
        quality.column_dimensions["B"].width = 34
        quality.column_dimensions["C"].width = 58
        quality.column_dimensions["D"].width = 24

        ws1 = wb.create_sheet(title="指标覆盖率矩阵")
        ws1.views.sheetView[0].showGridLines = True
        headers1 = ["序号", "模块分类", "核心指标名称", "来源表格", "覆盖状态", "关联报告章节"]
        ws1.append(headers1)
        for col_idx, _ in enumerate(headers1, 1):
            cell = ws1.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws1.row_dimensions[1].height = 26
        for idx, row in enumerate(coverage_data, 1):
            r_idx = idx + 1
            ws1.append([
                idx,
                row.get("category", ""),
                row.get("metric_name", ""),
                row.get("source_table", ""),
                row.get("status", "已纳入分析"),
                row.get("section_title", ""),
            ])
            for col_idx in range(1, len(headers1) + 1):
                cell = ws1.cell(row=r_idx, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", horizontal="center" if col_idx in [1, 5] else "left")
                if r_idx % 2 == 0:
                    cell.fill = zebra_fill

        ws2 = wb.create_sheet(title="全文数字对账总表")
        ws2.views.sheetView[0].showGridLines = True
        ws2.merge_cells("A1:I1")
        numeric_note = ws2["A1"]
        numeric_note.value = "说明：本表仅核对报告数值与原始单元格是否一致，不代表整份报告通过语义核验。"
        numeric_note.fill = warning_fill
        numeric_note.font = Font(name="Microsoft YaHei", size=10, color="92400E")
        numeric_note.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        numeric_note.border = thin_border
        ws2.row_dimensions[1].height = 30

        headers2 = [
            "序号", "报告章节", "报告所引数值", "对应指标绝对路径",
            "原始文件名称", "原始工作表(Sheet)", "单元格物理坐标", "原始单元格文本", "穿透核验状态",
        ]
        ws2.append(headers2)
        header_row = 2
        for col_idx, _ in enumerate(headers2, 1):
            cell = ws2.cell(row=header_row, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws2.row_dimensions[header_row].height = 26
        for idx, row in enumerate(reconciliation_data, 1):
            r_idx = idx + header_row
            ws2.append([
                idx,
                row.get("section_title", ""),
                row.get("cited_text", ""),
                row.get("metric_path", ""),
                row.get("file_name", ""),
                row.get("sheet_name", ""),
                row.get("cell_ref", ""),
                row.get("raw_value", ""),
                row.get("status", "已匹配"),
            ])
            for col_idx in range(1, len(headers2) + 1):
                cell = ws2.cell(row=r_idx, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", horizontal="center" if col_idx in [1, 3, 7, 9] else "left")
                if r_idx % 2 == 0:
                    cell.fill = zebra_fill
                if col_idx == 9:
                    cell.font = status_pass_font if cell.value == "100% 吻合" else warning_font

        for ws in [ws1, ws2]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    max_len = max(max_len, len(val_str.encode("gbk", errors="ignore")))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        clean_school = re.sub(r"[^a-zA-Z0-9_\u4e00-\u9fa5]", "", school_name) or "高校"
        prefix = "示例报告_" if is_demo else ""
        file_name = f"{prefix}数据穿透审计与指标覆盖清单_{clean_school}.xlsx"
        file_path = os.path.join(self.output_dir, file_name)
        wb.save(file_path)
        return file_path
