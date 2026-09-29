import os
import unittest
import openpyxl
from app.config import settings
from app.core.stc_parser import STCParser
from app.core.cell_lake import CellLake
from app.core.duckdb_engine import DuckDBEngine
from app.core.chart_service import ChartService
from app.core.audit_service import AuditService
from app.core.docx_exporter import DocxExporter
from app.core.excel_exporter import ExcelExporter
from app.pipeline.agent_runner import PiAgentRunner
from app.pipeline.jev_judge import JevJudge
from app.pipeline.stages import ReportPipeline

class TestXlsxAndCustomSchool(unittest.TestCase):
    def test_xlsx_composite_header_parsing(self):
        """验证 .xlsx 格式的跨行跨列复合表头与合并单元格前向填充"""
        test_file = os.path.join(settings.DATA_DIR, "test_composite.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "师资队伍结构"

        # 制作复合表头:
        # Row 1: [A1:A2 "学院名称"], [B1:C1 "高层次人才"], [D1:D2 "教师总数"]
        # Row 2: [B2 "领军人才"], [C2 "青年骨干"]
        ws["A1"] = "学院名称"
        ws.merge_cells("A1:A2")

        ws["B1"] = "高层次人才"
        ws["C1"] = "高层次人才"
        ws.merge_cells("B1:C1")

        ws["D1"] = "教师总数"
        ws.merge_cells("D1:D2")

        ws["B2"] = "领军人才"
        ws["C2"] = "青年骨干"

        # 数据行
        ws.append(["计算机学院", 8, 22, 120])
        ws.append(["数学学院", 5, 18, 95])
        wb.save(test_file)

        try:
            parsed = STCParser.parse_file(test_file)
            self.assertEqual(len(parsed), 1)
            sheet = parsed[0]
            print(f"[✓] 解析 .xlsx 复合表头路径: {sheet['header_paths']}")
            
            # 断言表头是否正确拼接
            self.assertEqual(sheet["header_paths"][0], "学院名称")
            self.assertIn("高层次人才 > 领军人才", sheet["header_paths"][1])
            self.assertIn("高层次人才 > 青年骨干", sheet["header_paths"][2])
            self.assertEqual(sheet["header_paths"][3], "教师总数")
            self.assertEqual(sheet["row_count"], 2)
            self.assertEqual(len(sheet["cells"]), 8)
        finally:
            if os.path.exists(test_file):
                os.remove(test_file)

    def test_custom_school_name_flow(self):
        """验证自定义高校名称（如北京大学）全流程联动"""
        cell_lake = CellLake(settings.SQLITE_PATH)
        duckdb_engine = DuckDBEngine(settings.DUCKDB_PATH)
        chart_service = ChartService(settings.CHARTS_DIR)
        pipeline = ReportPipeline(
            cell_lake=cell_lake,
            duckdb_engine=duckdb_engine,
            chart_service=chart_service,
            audit_service=AuditService(cell_lake),
            docx_exporter=DocxExporter(settings.EXPORTS_DIR),
            excel_exporter=ExcelExporter(settings.EXPORTS_DIR),
            agent_runner=PiAgentRunner(chart_service),
            jev_judge=JevJudge()
        )

        pipeline.set_custom_school_name("北京大学")
        self.assertEqual(pipeline.custom_school_name, "北京大学")

        outline = pipeline.plan_outline()
        self.assertGreaterEqual(len(outline), 1)

if __name__ == "__main__":
    unittest.main()
